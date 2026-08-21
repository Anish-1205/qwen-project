from __future__ import annotations

import csv
import json
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import load_workbook

from . import config
from .common import ToolError, resolve_file, validate_zip_archive

SUPPORTED = {".csv", ".xlsx"}


def _load(path: str, sheet: str | None, include_source: bool,
          row_limit: int, cell_limit: int) -> tuple[list[dict], list[str], int]:
    target = resolve_file(path, SUPPORTED, config.SPREADSHEET_MAX_FILE_BYTES)
    rows: list[dict] = []
    if target.suffix.lower() == ".csv":
        with target.open("r", encoding="utf-8-sig", newline="", errors="replace") as handle:
            reader = csv.DictReader(handle)
            headers = list(reader.fieldnames or [])
            if any(not str(header).strip() for header in headers):
                raise ToolError("invalid_columns", "CSV column names must not be empty.")
            if len(set(headers)) != len(headers):
                raise ToolError("duplicate_columns", "CSV column names must be unique.")
            if include_source and "_source_file" in headers:
                raise ToolError("column_conflict", "Input already contains the reserved _source_file column.")
            cell_count = len(headers) + (1 if include_source else 0)
            if cell_count > cell_limit:
                raise ToolError("cell_limit", "Spreadsheet cell limit exceeded.")
            for row in reader:
                row_cells = sum(len(value) if key is None and isinstance(value, list) else 1 for key, value in row.items())
                row_cells += 1 if include_source else 0
                if cell_count + row_cells > cell_limit:
                    raise ToolError("cell_limit", "Spreadsheet cell limit exceeded.")
                if len(rows) >= row_limit:
                    raise ToolError("row_limit", "Spreadsheet row limit exceeded.")
                cell_count += row_cells
                rows.append(dict(row))
    else:
        validate_zip_archive(target)
        workbook = load_workbook(target, read_only=True, data_only=True)
        if sheet is not None and sheet not in workbook.sheetnames:
            workbook.close()
            raise ToolError("sheet_not_found", "The requested worksheet does not exist.", {"sheet": sheet})
        worksheet = workbook[sheet] if sheet else workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, None)
        headers = [str(value) if value is not None else f"column_{i + 1}" for i, value in enumerate(first or [])]
        if len(set(headers)) != len(headers):
            workbook.close()
            raise ToolError("duplicate_columns", "XLSX column names must be unique.")
        if include_source and "_source_file" in headers:
            workbook.close()
            raise ToolError("column_conflict", "Input already contains the reserved _source_file column.")
        if max(worksheet.max_row - 1, 0) > row_limit:
            workbook.close()
            raise ToolError("row_limit", "Spreadsheet row limit exceeded.")
        cell_count = (len(headers) + (1 if include_source else 0)) * max(worksheet.max_row, 0)
        if cell_count > cell_limit:
            workbook.close()
            raise ToolError("cell_limit", "Spreadsheet cell limit exceeded.")
        for values in iterator:
            rows.append(dict(zip(headers, values)))
        workbook.close()
    if include_source:
        headers.append("_source_file")
        for row in rows:
            row["_source_file"] = str(target)
    return rows, headers, cell_count


def _require_columns(columns: list[str], headers: list[str]) -> None:
    missing = [column for column in columns if column not in headers]
    if missing:
        raise ToolError("column_not_found", "One or more columns do not exist.", {"columns": missing})


def _number(value: Any, column: str, error_code: str = "nonnumeric_value") -> float:
    if type(value) in (int, float):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ToolError(error_code, f"Column '{column}' contains a nonnumeric value.") from exc


def analyze_spreadsheet(paths: list[str], operations: list[dict], sheet: str | None = None,
                        include_source: bool = False) -> dict:
    combined: list[dict] = []
    headers: list[str] | None = None
    combined_cells = 0
    for path in paths:
        rows, current_headers, current_cells = _load(
            path,
            sheet,
            include_source,
            config.SPREADSHEET_MAX_ROWS - len(combined),
            config.SPREADSHEET_MAX_CELLS - combined_cells,
        )
        combined_cells += current_cells
        if combined_cells > config.SPREADSHEET_MAX_CELLS:
            raise ToolError("cell_limit", "Combined spreadsheet cell limit exceeded.")
        if headers is None:
            headers = current_headers
        elif headers != current_headers:
            raise ToolError("incompatible_schemas", "Input files do not have compatible columns.")
        combined.extend(rows)
        if len(combined) > config.SPREADSHEET_MAX_ROWS:
            raise ToolError("row_limit", "Combined spreadsheet row limit exceeded.")
    headers = headers or []
    rows = combined
    group_columns: list[str] = []
    for operation in operations:
        op = operation["op"]
        if op == "select":
            columns = operation["columns"]
            _require_columns(columns, headers)
            if group_columns and any(column not in columns for column in group_columns):
                raise ToolError("invalid_operation", "A select operation cannot remove active group-by columns.", {"group_by": group_columns})
            rows = [{column: row.get(column) for column in columns} for row in rows]
            headers = columns
        elif op == "filter":
            column, operator_name, expected = operation["column"], operation["operator"], operation["value"]
            _require_columns([column], headers)
            def matches(row: dict) -> bool:
                actual = row.get(column)
                if operator_name in {"gt", "gte", "lt", "lte"}:
                    left, right = _number(actual, column), _number(expected, column)
                    return {"gt": left > right, "gte": left >= right, "lt": left < right, "lte": left <= right}[operator_name]
                if operator_name == "contains": return str(expected).casefold() in str(actual).casefold()
                if operator_name == "in": return actual in expected
                if type(actual) in (int, float) or type(expected) in (int, float):
                    equal = _number(actual, column) == _number(expected, column)
                else:
                    equal = actual == expected
                return equal if operator_name == "eq" else not equal
            rows = [row for row in rows if matches(row)]
        elif op == "group_by":
            group_columns = operation["columns"]
            _require_columns(group_columns, headers)
        elif op == "aggregate":
            column, function = operation.get("column"), operation["function"]
            if function != "count":
                _require_columns([column], headers)
            groups: dict[tuple, list[dict]] = {(): rows} if not group_columns else {}
            if group_columns:
                for row in rows:
                    groups.setdefault(tuple(row.get(c) for c in group_columns), []).append(row)
            output = []
            for key, members in groups.items():
                result = len(members) if function == "count" else {
                    "sum": sum, "min": min, "max": max, "mean": mean
                }[function]([_number(item.get(column), column, "nonnumeric_aggregation") for item in members])
                output.append({**dict(zip(group_columns, key)), operation.get("as") or f"{function}_{column or 'rows'}": result})
            rows = output
            headers = list(rows[0]) if rows else group_columns + [operation.get("as") or f"{function}_{column or 'rows'}"]
        elif op == "sort":
            column = operation["column"]
            _require_columns([column], headers)
            rows.sort(key=lambda row: (row.get(column) is None, row.get(column)), reverse=operation.get("direction", "asc") == "desc")
        elif op == "limit":
            rows = rows[:operation["count"]]
    returned: list[dict] = []
    for row in rows[:config.SPREADSHEET_MAX_RETURNED_ROWS]:
        if len(json.dumps(returned + [row], ensure_ascii=False, default=str)) > config.MAX_RESULT_CHARS:
            break
        returned.append(row)
    result = {"columns": list(headers), "rows": returned, "row_count": len(rows), "returned_rows": len(returned),
              "truncated": len(rows) > len(returned), "source_count": len(paths)}
    while len(json.dumps(result, ensure_ascii=False, default=str)) > config.MAX_RESULT_CHARS:
        result["truncated"] = True
        if result["rows"]:
            result["rows"].pop()
            result["returned_rows"] = len(result["rows"])
        elif len(result["columns"]) > 1:
            result["columns"] = result["columns"][: max(1, len(result["columns"]) // 2)]
            result["columns_truncated"] = True
        else:
            result["columns"] = []
            result["columns_truncated"] = True
            break
    return result
