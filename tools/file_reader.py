from __future__ import annotations

import csv
import io
import json

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from . import config
from .common import ToolError, resolve_file, truncate_text, validate_zip_archive

SUPPORTED = {".txt", ".csv", ".json", ".docx", ".pdf", ".xlsx"}


def _decode(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replacement"


def read_file(path: str) -> dict:
    target = resolve_file(path, SUPPORTED, config.MAX_FILE_BYTES)
    suffix = target.suffix.lower()
    metadata: dict = {"path": str(target), "format": suffix[1:], "bytes": target.stat().st_size}
    extraction_truncated = False
    try:
        if suffix == ".txt":
            text, encoding = _decode(target.read_bytes())
            metadata["encoding"] = encoding
        elif suffix == ".csv":
            raw, encoding = _decode(target.read_bytes())
            reader = csv.reader(io.StringIO(raw))
            rows, total = [], 0
            for row in reader:
                total += 1
                if len(rows) < config.FILE_CSV_PREVIEW_ROWS:
                    rows.append(row)
                else:
                    extraction_truncated = True
            metadata.update({"encoding": encoding, "headers": rows[0] if rows else [], "row_count": max(total - 1, 0)})
            text = "\n".join(", ".join(str(cell) for cell in row) for row in rows)
        elif suffix == ".json":
            raw, encoding = _decode(target.read_bytes())
            payload = json.loads(raw)
            metadata.update({"encoding": encoding, "root_type": type(payload).__name__})
            text = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
        elif suffix == ".docx":
            validate_zip_archive(target)
            document = Document(target)
            lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            metadata.update({"paragraphs": len(document.paragraphs), "tables": len(document.tables)})
            text = "\n".join(lines)
        elif suffix == ".pdf":
            reader = PdfReader(str(target))
            pages = [(page.extract_text() or "") for page in reader.pages]
            metadata["pages"] = len(pages)
            text = "\n\n".join(f"Page {index + 1}\n{value}" for index, value in enumerate(pages))
        else:
            validate_zip_archive(target)
            workbook = load_workbook(target, read_only=True, data_only=True)
            previews: list[str] = []
            sheets: list[dict] = []
            for sheet in workbook.worksheets:
                rows = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= config.FILE_XLSX_PREVIEW_ROWS:
                        extraction_truncated = True
                        break
                    rows.append([cell for cell in row])
                sheets.append({"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column})
                previews.append(f"Sheet: {sheet.title}\n" + "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in rows))
            workbook.close()
            metadata["sheets"] = sheets
            text = "\n\n".join(previews)
    except ToolError:
        raise
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        raise ToolError("file_read_error", "The file could not be parsed.", {"reason": str(exc)[:300]}) from exc
    text, truncated = truncate_text(text)
    metadata.update({"characters": len(text), "truncated": truncated or extraction_truncated})
    return {"text": text, "metadata": metadata}
