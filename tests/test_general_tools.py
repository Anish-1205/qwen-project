from __future__ import annotations

import csv
import json
from pathlib import Path
from unittest.mock import Mock, patch

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from intent_classifier import DeterministicIntentRouter, IntentClassifier, IntentDecision
from orchestrator import ConversationOrchestrator, DEFAULT_SYSTEM_PROMPT
from logging_utils import sanitize_tool_log_payload
from tools import ToolCall, ToolManager
from tools.directory_listing import list_directory
from tools.common import validate_public_url
from tools.config import _bounded_int
from tools.file_reader import read_file
from tools.spreadsheet import analyze_spreadsheet
from tools.weather import FORECAST_URL, GEOCODING_URL, weather
from tools.web_fetch import fetch_webpage


def _xlsx(path: Path, rows, title="Data"):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows: sheet.append(row)
    book.save(path)


def test_calculator_modes_and_failures():
    manager = ToolManager()
    assert manager.execute(ToolCall("calculator", {"expression": "(2 + 3) ** 2"})).data["value"] == 25
    assert manager.execute(ToolCall("calculator", {"aggregate": "mean", "values": [2, 4, 6]})).data["value"] == 4
    assert manager.execute(ToolCall("calculator", {"expression": "1", "aggregate": "sum", "values": [1]})).error_details["code"] == "validation_error"
    assert manager.execute(ToolCall("calculator", {"expression": "(-1) ** 0.5"})).error_details["code"] == "numeric_domain_error"


def test_parser_contains_malformed_tagged_calls_and_ignores_normal_json():
    manager = ToolManager()
    malformed = manager.parse_tool_calls('<tool_call>{"arguments":{"expression":"1+1"}}</tool_call>')
    assert len(malformed) == 1 and malformed[0].name == "__invalid_tool_call__"
    assert manager.parse_tool_calls('{"name":"Pune","temperature":25}') == []
    nonfinite = manager.parse_tool_calls('<tool_call>{"name":"calculator","arguments":{"aggregate":"count","values":[NaN]}}</tool_call>')
    assert nonfinite[0].name == "__invalid_tool_call__"
    assert manager.parse_tool_calls('<tool_call>{"name":"calculator"')[0].name == "__invalid_tool_call__"
    partial_batch = manager.parse_tool_calls('<tool_call>{"name":"calculator","arguments":{"expression":"1+1"}}</tool_call><tool_call>{')
    assert [call.name for call in partial_batch] == ["calculator", "__invalid_tool_call__"]


def test_dispatcher_bounds_the_complete_result_envelope(tmp_path):
    wide = tmp_path / "wide.csv"
    wide.write_text(",".join(f"column_{index}" for index in range(10_000)) + "\n", encoding="utf-8")
    result = ToolManager().execute(ToolCall("read_file", {"path": str(wide)}))
    assert result.ok and len(json.dumps(result.payload(), ensure_ascii=False)) <= 20_000
    assert result.data["metadata"]["truncated"] is True


def test_directory_listing_success_and_failure(tmp_path):
    (tmp_path / "b.txt").write_text("b", encoding="utf-8")
    (tmp_path / "a.csv").write_text("a\n1", encoding="utf-8")
    result = list_directory(str(tmp_path), extensions=["txt"])
    assert [entry["name"] for entry in result["entries"]] == ["b.txt"]
    with pytest.raises(Exception) as error: list_directory(str(tmp_path / "missing"))
    assert getattr(error.value, "code", None) == "path_not_found"
    nested = tmp_path / "nested"; nested.mkdir(); (nested / "inside.txt").write_text("x", encoding="utf-8")
    recursive = list_directory(str(tmp_path), recursive=True)
    assert "inside.txt" in [entry["name"] for entry in recursive["entries"]]


@pytest.mark.parametrize("extension", ["txt", "csv", "json", "docx", "pdf", "xlsx"])
def test_file_reader_formats(tmp_path, extension):
    path = tmp_path / f"sample.{extension}"
    if extension == "txt": path.write_text("hello text", encoding="utf-8")
    elif extension == "csv": path.write_text("name,value\na,2\n", encoding="utf-8")
    elif extension == "json": path.write_text('{"hello":"world"}', encoding="utf-8")
    elif extension == "docx":
        document = Document(); document.add_paragraph("hello docx"); document.save(path)
    elif extension == "pdf":
        writer = PdfWriter(); writer.add_blank_page(width=72, height=72)
        with path.open("wb") as handle: writer.write(handle)
    else: _xlsx(path, [["name", "value"], ["a", 2]])
    result = read_file(str(path))
    assert result["metadata"]["format"] == extension
    assert "truncated" in result["metadata"]


def test_file_reader_failure(tmp_path):
    path = tmp_path / "bad.bin"; path.write_bytes(b"data")
    with pytest.raises(Exception) as error: read_file(str(path))
    assert error.value.code == "unsupported_format"


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_spreadsheet_filter_group_aggregate(tmp_path, extension):
    path = tmp_path / f"sales.{extension}"
    rows = [["region", "revenue"], ["east", 10], ["west", 5], ["east", 20]]
    if extension == "csv":
        with path.open("w", newline="", encoding="utf-8") as handle: csv.writer(handle).writerows(rows)
    else: _xlsx(path, rows)
    result = analyze_spreadsheet([str(path)], [
        {"op": "filter", "column": "revenue", "operator": "gte", "value": 10},
        {"op": "group_by", "columns": ["region"]},
        {"op": "aggregate", "column": "revenue", "function": "sum", "as": "total"},
    ])
    assert result["rows"] == [{"region": "east", "total": 30.0}]


def test_spreadsheet_specific_failure(tmp_path):
    path = tmp_path / "data.csv"; path.write_text("a\n1\n", encoding="utf-8")
    with pytest.raises(Exception) as error: analyze_spreadsheet([str(path)], [{"op": "select", "columns": ["missing"]}])
    assert error.value.code == "column_not_found"
    with patch("tools.spreadsheet.config.SPREADSHEET_MAX_CELLS", 1), pytest.raises(Exception) as limit_error:
        analyze_spreadsheet([str(path)], [])
    assert limit_error.value.code == "cell_limit"
    duplicate = tmp_path / "duplicate.csv"; duplicate.write_text("a,a\n1,2\n", encoding="utf-8")
    with pytest.raises(Exception) as duplicate_error:
        analyze_spreadsheet([str(duplicate)], [])
    assert duplicate_error.value.code == "duplicate_columns"


def test_spreadsheet_xlsx_and_combined_cell_limits(tmp_path):
    duplicate = tmp_path / "duplicate.xlsx"
    _xlsx(duplicate, [["a", "a"], [1, 2]])
    with pytest.raises(Exception) as duplicate_error:
        analyze_spreadsheet([str(duplicate)], [])
    assert duplicate_error.value.code == "duplicate_columns"

    source = tmp_path / "source.csv"
    source.write_text("a\n1\n", encoding="utf-8")
    with patch("tools.spreadsheet.config.SPREADSHEET_MAX_CELLS", 3), pytest.raises(Exception) as source_error:
        analyze_spreadsheet([str(source)], [], include_source=True)
    assert source_error.value.code == "cell_limit"

    first = tmp_path / "first.csv"; first.write_text("a\n", encoding="utf-8")
    second = tmp_path / "second.csv"; second.write_text("a\n", encoding="utf-8")
    with patch("tools.spreadsheet.config.SPREADSHEET_MAX_CELLS", 1), pytest.raises(Exception) as combined_error:
        analyze_spreadsheet([str(first), str(second)], [])
    assert combined_error.value.code == "cell_limit"


def test_spreadsheet_file_and_row_limits(tmp_path):
    path = tmp_path / "rows.csv"
    path.write_text("a\n1\n2\n", encoding="utf-8")
    with patch("tools.spreadsheet.config.SPREADSHEET_MAX_ROWS", 1), pytest.raises(Exception) as row_error:
        analyze_spreadsheet([str(path)], [])
    assert row_error.value.code == "row_limit"
    with patch("tools.spreadsheet.config.SPREADSHEET_MAX_FILE_BYTES", 1), pytest.raises(Exception) as file_error:
        analyze_spreadsheet([str(path)], [])
    assert file_error.value.code == "file_too_large"
    first = tmp_path / "first.csv"; first.write_text("a\n1\n", encoding="utf-8")
    second = tmp_path / "second.csv"; second.write_text("a\n2\n", encoding="utf-8")
    with patch("tools.spreadsheet.config.SPREADSHEET_MAX_ROWS", 1), pytest.raises(Exception) as combined_row_error:
        analyze_spreadsheet([str(first), str(second)], [])
    assert combined_row_error.value.code == "row_limit"


def test_spreadsheet_numeric_equality_and_empty_count(tmp_path):
    path = tmp_path / "values.csv"; path.write_text("value\n10\n20\n", encoding="utf-8")
    filtered = analyze_spreadsheet([str(path)], [{"op": "filter", "column": "value", "operator": "eq", "value": 10}])
    assert filtered["row_count"] == 1
    empty = tmp_path / "empty.csv"; empty.write_text("value\n", encoding="utf-8")
    counted = analyze_spreadsheet([str(empty)], [{"op": "aggregate", "function": "count", "as": "rows"}])
    assert counted["rows"] == [{"rows": 0}]
    with pytest.raises(Exception) as operation_error:
        analyze_spreadsheet([str(path)], [{"op": "group_by", "columns": ["value"]}, {"op": "select", "columns": []}])
    assert operation_error.value.code in {"column_not_found", "invalid_operation"}


def test_web_fetch_success_and_failure():
    response = Mock()
    response.is_redirect = response.is_permanent_redirect = False
    response.headers = {"content-type": "text/html; charset=utf-8"}
    response.encoding = "utf-8"
    response.iter_content.return_value = [b"<html><head><title>T</title><script>x</script></head><body><nav>N</nav>Hello world</body></html>"]
    response.raise_for_status.return_value = None
    with patch("tools.web_fetch.validate_public_url", side_effect=lambda value: value), patch("tools.web_fetch.requests.Session") as session:
        session.return_value.get.return_value = response
        result = fetch_webpage("https://example.com")
    assert result["title"] == "T" and result["text"] == "Hello world"
    with pytest.raises(Exception) as error: fetch_webpage("file:///etc/passwd")
    assert error.value.code == "invalid_url"


@pytest.mark.parametrize("url", ["https://user:pass@example.com/", "https://example.com:bad/path"])
def test_web_fetch_rejects_malformed_or_credentialed_urls(url):
    result = ToolManager().execute(ToolCall("fetch_webpage", {"url": url}))
    assert not result.ok and result.error_details["code"] == "invalid_url"


def test_web_fetch_revalidates_and_blocks_redirect_targets():
    response = Mock()
    response.is_redirect = True; response.is_permanent_redirect = False
    response.headers = {"location": "http://internal.test/secret"}
    def addresses(host, *args, **kwargs):
        address = "127.0.0.1" if host == "internal.test" else "93.184.216.34"
        return [(2, 1, 6, "", (address, 80))]
    with patch("tools.common.socket.getaddrinfo", side_effect=addresses), patch("tools.web_fetch.requests.Session") as session:
        session.return_value.get.return_value = response
        with pytest.raises(Exception) as error:
            fetch_webpage("https://example.com/start")
    assert error.value.code == "blocked_host"
    response.close.assert_called()


def test_web_fetch_rejects_redirect_without_location_and_multicast():
    response = Mock(status_code=302)
    response.is_redirect = response.is_permanent_redirect = False
    response.headers = {}
    with patch("tools.web_fetch.validate_public_url", side_effect=lambda value: value), patch("tools.web_fetch.requests.Session") as session:
        session.return_value.get.return_value = response
        with pytest.raises(Exception) as error:
            fetch_webpage("https://example.com/start")
    assert error.value.code == "invalid_redirect"
    with patch("tools.common.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("224.0.0.1", 80))]):
        blocked = ToolManager().execute(ToolCall("fetch_webpage", {"url": "http://multicast.example/"}))
    assert not blocked.ok and blocked.error_details["code"] == "blocked_host"


def test_url_validation_rejects_controls_site_local_and_oversized_redirects():
    for url in ("https://exa\nmple.com/path", "https://example.com/\x00path", "https://example.com/\x7f"):
        with pytest.raises(Exception) as error:
            validate_public_url(url)
        assert error.value.code == "invalid_url"
    with patch("tools.common.socket.getaddrinfo", return_value=[(10, 1, 6, "", ("fec0::1", 80, 0, 0))]):
        with pytest.raises(Exception) as error:
            validate_public_url("http://site-local.example/")
    assert error.value.code == "blocked_host"

    response = Mock(status_code=302)
    response.is_redirect = True; response.is_permanent_redirect = False
    response.headers = {"location": "https://example.com/" + "x" * 5_000}
    with patch("tools.common.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 80))]), patch("tools.web_fetch.requests.Session") as session:
        session.return_value.get.return_value = response
        with pytest.raises(Exception) as error:
            fetch_webpage("https://example.com/start")
    assert error.value.code == "invalid_url"


def test_integer_environment_bounds_handle_arbitrarily_large_values(monkeypatch):
    monkeypatch.setenv("TOOLS_TEST_HUGE_INT", "9" * 4_000)
    assert _bounded_int("TOOLS_TEST_HUGE_INT", 7, 1, 100) == 100


def test_weather_place_success_and_coordinate_failure():
    geocode = Mock(); geocode.raise_for_status.return_value = None
    geocode.json.return_value = {"results": [{"name": "Pune", "country": "India", "latitude": 18.52, "longitude": 73.85, "timezone": "Asia/Kolkata"}]}
    forecast = Mock(); forecast.raise_for_status.return_value = None
    forecast.json.return_value = {"timezone": "Asia/Kolkata", "current": {"time": "now", "temperature_2m": 25, "apparent_temperature": 26,
        "precipitation": 0, "weather_code": 0, "wind_speed_10m": 5, "wind_direction_10m": 90},
        "daily": {"time": ["d1"], "weather_code": [0], "temperature_2m_max": [30], "temperature_2m_min": [20], "precipitation_probability_max": [5]}}
    def fake_get(url, **kwargs): return geocode if url == GEOCODING_URL else forecast
    with patch("tools.weather.requests.get", side_effect=fake_get): result = weather(place="Pune", forecast_days=1)
    assert result["current"]["condition"] == "Clear sky"
    with pytest.raises(Exception) as error: weather(latitude=100, longitude=0)
    assert error.value.code == "invalid_coordinates"


def test_weather_incomplete_upstream_data_is_structured():
    response = Mock(); response.raise_for_status.return_value = None
    response.json.return_value = {"current": {"temperature_2m": 20}, "daily": {"time": ["d1"]}}
    with patch("tools.weather.requests.get", return_value=response), pytest.raises(Exception) as error:
        weather(latitude=10, longitude=10, forecast_days=1)
    assert error.value.code == "incomplete_weather_data"


def test_routing_tool_flags_semantic_repair_and_mixed():
    router = DeterministicIntentRouter()
    assert router.analyze("Fetch https://example.com and read it", []).tool_use is True
    assert router.analyze("How does weather forecasting work?", []).tool_use is False
    assert router.analyze("Calculate the sum of 10, 20, and 30", []).tool_use is True
    directory = router.analyze("List the files in this folder", [])
    assert directory.tool_use is True and directory.document_read is False and directory.complete
    for request in ("What's 2+2?", "Weather Pune", "What's Pune's temperature?", "Pick a number between 1 and 10", r"What does C:\notes.txt say?"):
        assert router.analyze(request, []).tool_use is not False
    for request in ("What files are here?", "Which files are in this folder?", "What folders are here?", r"Which directories are in C:\Data?", "List data/", "List ~/Downloads", "Show data/reports", "Find in reports/2026", "List common files in this folder", "Show standard files in the current directory", r"List files normally found in C:\project", r"Which standard files are in C:\project?", "Extract text from report.pdf", "Inspect report.pdf", "Parse data.json", "Load report.xlsx", "In sales.csv, group rows by region", "In sales.csv, mean amount", "Limit sales.csv to 10 rows"):
        decision = router.analyze(request, [])
        assert decision.tool_use is True and decision.document_read is False
    for discussion in ("How do current weather APIs work?", "Write Python code to fetch https://example.com", "Do not fetch https://example.com", "Show me how files and directories work", "Explain how to filter a spreadsheet", "Teach me how to group rows in a CSV", "Show code to filter sales.csv", "How can I filter sales.csv?", "How should I group sales.csv?", "What filter should I use on sales.csv?", "Can you explain how to filter sales.csv?", "Could you show me how to list files?", "Should I sort sales.csv by date?", "Do not group sales.csv by region", "Never aggregate sales.csv", "Do not limit sales.csv"):
        assert router.analyze(discussion, []).tool_use is False
    for discussion in ("What files are in a typical Python package?", "Which directories are in the Unix filesystem hierarchy?", "List common files in a Python package", "Show the files normally found in a web project"):
        decision = router.analyze(discussion, [])
        assert decision.tool_use is False and decision.document_read is False
    assert router.analyze("Remember I prefer metric units and get current weather in Pune", []).memory_write is True
    assert router.analyze("Remember I prefer metric units and get current weather in Pune", []).tool_use is True
    mixed_document = router.analyze("What does the travel policy say, and get current weather in Pune?", [])
    assert mixed_document.document_read is True and mixed_document.tool_use is True
    outputs = iter(["bad", '{"memory_read":false,"memory_write":false,"document_read":false,"tool_use":true,"general_chat":false}'])
    classifier = IntentClassifier(lambda messages, **kwargs: next(outputs), logger=lambda _: None)
    assert classifier.classify("Handle this utility request", []).tool_use is True
    broken = IntentClassifier(lambda messages, **kwargs: "bad", logger=lambda _: None)
    assert broken.classify("Handle this", []).tool_use is False


def test_tool_log_sanitizer_redacts_and_bounds():
    rendered = sanitize_tool_log_payload({"api_key": "secret", "url": "https://user:pass@example.com/a?q=sensitive", "text": "x" * 5000}, limit=300)
    assert "secret" not in rendered and "pass" not in rendered and "sensitive" not in rendered
    assert "5000 characters" in rendered and len(rendered) <= 300
    bounded = sanitize_tool_log_payload({"value": "x" * 1000}, limit=100)
    assert len(bounded) == 100 and bounded.endswith("[truncated]")
    malformed = sanitize_tool_log_payload({"url": "https://user:pass@example.com:bad/path?q=secret"})
    assert "pass" not in malformed and "secret" not in malformed
    embedded = sanitize_tool_log_payload({"passwd": "hidden", "message": "see HTTPS://user:pass@example.com/a?token=hidden"})
    assert "hidden" not in embedded and "user:pass" not in embedded
    assert len(sanitize_tool_log_payload({"value": "long" * 20}, limit=5)) == 5
    malformed_url = sanitize_tool_log_payload({"url": " https://user:pass@/path?q=hidden"})
    assert "user" not in malformed_url and "pass" not in malformed_url and "hidden" not in malformed_url
    opaque_query = sanitize_tool_log_payload({"url": "https://example.com/path?bare-secret-token"})
    assert "bare-secret-token" not in opaque_query


class FakeTokenizer:
    eos_token_id = 0
    def apply_chat_template(self, messages, **kwargs): return " ".join(str(m.get("content", "")) for m in messages)
    def __call__(self, prompt, *args, **kwargs): return {"input_ids": list(range(len(prompt.split())))}

class FakeMemory:
    last_retrieval_stats = {"facts": []}

class QueueOrchestrator(ConversationOrchestrator):
    def __init__(self, outputs, **kwargs): self.outputs=list(outputs); self.inputs=[]; self.overrides=[]; super().__init__(FakeTokenizer(), object(), FakeMemory(), logger=lambda _: None, **kwargs)
    def generate_reply(self, messages, **overrides):
        self.inputs.append([dict(m) for m in messages]); self.overrides.append(overrides)
        return self.outputs.pop(0)


def test_multi_tool_loop_budget_and_ephemeral_history():
    calls = ['<tool_call>{"name":"calculator","arguments":{"expression":"1+1"}}</tool_call>',
             '<tool_call>{"name":"calculator","arguments":{"expression":"2+2"}}</tool_call>', "done"]
    orchestrator = QueueOrchestrator(calls)
    history, reply, _, _ = orchestrator.process_turn("Calculate 1+1 and 2+2", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "done" and len(orchestrator.last_tool_executions) == 2
    assert "untrusted data" in orchestrator.inputs[0][0]["content"]
    assert [m["role"] for m in history] == ["system", "user", "assistant"]
    tool_ids = [m["tool_call_id"] for m in orchestrator.inputs[-1] if m["role"] == "tool"]
    assert tool_ids == ["tool_call_1", "tool_call_2"]
    exhausted = QueueOrchestrator([calls[0], "budget answer"])
    with patch("orchestrator.MAX_TOOL_CALLS_PER_TURN", 1):
        _, reply, _, _ = exhausted.process_turn("Calculate 1+1 and 2+2", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "budget answer" and exhausted.overrides[-1] == {}
    uncooperative = QueueOrchestrator([calls[0], calls[1]])
    with patch("orchestrator.MAX_TOOL_CALLS_PER_TURN", 1):
        _, reply, _, _ = uncooperative.process_turn("Calculate 1+1 and 2+2", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert "tool-call limit" in reply and "<tool_call>" not in reply


def test_multiple_calls_in_one_model_message_execute_in_order():
    combined = ('<tool_call>{"name":"calculator","arguments":{"expression":"3+4"}}</tool_call>'
                '<tool_call>{"name":"calculator","arguments":{"expression":"5+6"}}</tool_call>')
    orchestrator = QueueOrchestrator([combined, "results ready"])
    _, reply, _, _ = orchestrator.process_turn("Calculate 3+4 and 5+6", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "results ready"
    assert [result.data["value"] for result in orchestrator.last_tool_executions] == [7, 11]


def test_tool_call_ids_are_unique_and_malformed_calls_can_recover():
    combined = ('<tool_call>{"id":"tool_call_2","name":"calculator","arguments":{"expression":"1+1"}}</tool_call>'
                '<tool_call>{"name":"calculator","arguments":{"expression":"2+2"}}</tool_call>')
    orchestrator = QueueOrchestrator([combined, "done"])
    _, reply, _, _ = orchestrator.process_turn("Calculate both values", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "done"
    tool_ids = [message["tool_call_id"] for message in orchestrator.inputs[-1] if message["role"] == "tool"]
    assistant_ids = [message["tool_calls"][0]["id"] for message in orchestrator.inputs[-1] if message["role"] == "assistant" and message.get("tool_calls")]
    assert tool_ids == assistant_ids == ["tool_call_2", "tool_call_2_2"]

    malformed = QueueOrchestrator(['<tool_call>{"name":"calculator"', "recovered"])
    _, reply, _, _ = malformed.process_turn("Calculate this", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "recovered" and len(malformed.last_tool_executions) == 1
    assert malformed.last_tool_executions[0].error_details["code"] == "unknown_tool"


def test_tool_payload_context_serialization_stays_valid_and_bounded():
    payload = {"ok": True, "tool": "read_file", "data": {"text": "x" * 10_000}, "meta": {}}
    rendered = ConversationOrchestrator._serialize_tool_payload(payload, 500)
    assert len(rendered) <= 500 and json.loads(rendered)["meta"]["context_truncated"] is True
    minimal = ConversationOrchestrator._serialize_tool_payload({"ok": True, "tool": "x" * 1_000, "data": {"text": "x" * 1_000}}, 100)
    assert len(minimal) <= 100 and json.loads(minimal)["ok"] is True


def test_cumulative_tool_arguments_and_results_stay_within_context_budget():
    expression = "1+" * 5_000 + "1"
    call = f'<tool_call>{{"name":"calculator","arguments":{{"expression":{json.dumps(expression)}}}}}</tool_call>'
    orchestrator = QueueOrchestrator([call] * 8 + ["budget answer"])
    with patch("orchestrator.MAX_TOOL_CALLS_PER_TURN", 8), patch("orchestrator.MAX_TOOL_CONTEXT_CHARS", 2_000):
        _, reply, _, _ = orchestrator.process_turn("Calculate these values", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "budget answer"
    tool_messages = [message for message in orchestrator.inputs[-1] if message["role"] == "tool"]
    assistant_calls = [message["tool_calls"][0] for message in orchestrator.inputs[-1] if message["role"] == "assistant" and message.get("tool_calls")]
    context_chars = sum(len(message["content"]) for message in tool_messages)
    context_chars += sum(len(json.dumps(call["function"]["arguments"], ensure_ascii=True)) for call in assistant_calls)
    assert context_chars <= 2_000 and all(message["content"] != "{}" for message in tool_messages)


def test_end_to_end_three_spreadsheet_pipeline(tmp_path):
    paths=[]
    for index, values in enumerate(([10, 5], [20], [7, 8, 10])):
        path=tmp_path / f"sales{index}.xlsx"; _xlsx(path, [["revenue"], *[[v] for v in values]]); paths.append(path)
    outputs=[f'<tool_call>{{"name":"list_directory","arguments":{{"path":{json.dumps(str(tmp_path))},"extensions":["xlsx"]}}}}</tool_call>']
    for path in paths:
        outputs.append(f'<tool_call>{{"name":"analyze_spreadsheet","arguments":{{"paths":[{json.dumps(str(path))}],"operations":[{{"op":"aggregate","function":"sum","column":"revenue","as":"total"}}]}}}}</tool_call>')
    outputs += ['<tool_call>{"name":"calculator","arguments":{"aggregate":"sum","values":[15,20,25]}}</tool_call>', "Total revenue is 60."]
    orchestrator=QueueOrchestrator(outputs)
    _, reply, _, _=orchestrator.process_turn(f"List the spreadsheets in {tmp_path} and give me total revenue", [{"role":"system","content":DEFAULT_SYSTEM_PROMPT}], turn_number=1)
    assert reply == "Total revenue is 60." and [r.call.name for r in orchestrator.last_tool_executions] == ["list_directory", "analyze_spreadsheet", "analyze_spreadsheet", "analyze_spreadsheet", "calculator"]
    assert orchestrator.last_tool_executions[-1].data["value"] == 60
